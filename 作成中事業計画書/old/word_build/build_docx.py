# -*- coding: utf-8 -*-
"""章データ（chapters/row*.json）と前段データから、参考様式の docx を組み立てる。
使い方: python build_docx.py <出力docx> [--template <テンプレdocx>]
- python-docx のAPIで要素を作り、APIが無い要素（塗り・枠線・固定レイアウト等）は
  Wordの規格の並び順どおりに挿入する（並び順違反でWordが開けない事故の防止）。"""
import os, re, sys, json, copy, glob
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_COLOR_INDEX
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8')
HERE = os.path.dirname(os.path.abspath(__file__))
FONT = 'BIZ UDPゴシック'
NUM_FONT = 'BIZ UDゴシック'   # 表の数字列（桁がそろう等幅書体）
C_CH = '2E5496'      # 章見出し
C_H2 = '1F497D'      # 小見出し
C_RED = 'C00000'
C_BLUE = '0070C0'
C_GRAY = '404040'
C_NOTE = '404040'
HEADER_FILL = '2E5496'
BORDER = '7F7F7F'

# ---------------------------------------------------------------- 規格順の挿入ヘルパー
TCPR_ORDER = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark']
TBLPR_ORDER = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription']


def _insert_ordered(parent, child, order):
    name = child.tag.split('}')[1]
    for old in parent.findall(qn('w:' + name)):
        parent.remove(old)
    idx = order.index(name)
    for i, existing in enumerate(list(parent)):
        en = existing.tag.split('}')[1]
        if en in order and order.index(en) > idx:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


def set_cell_fill(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    _insert_ordered(tcPr, shd, TCPR_ORDER)


def set_cell_margins(cell, top=0, bottom=0, left=57, right=57):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for side, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    _insert_ordered(tcPr, mar, TCPR_ORDER)


def set_table_props(table, widths_cm, borders=True, cell_mar=(0, 57, 0, 57), align_center=False):
    tbl = table._tbl
    tblPr = tbl.tblPr
    total = int(sum(widths_cm) * 567)
    tblW = OxmlElement('w:tblW'); tblW.set(qn('w:w'), str(total)); tblW.set(qn('w:type'), 'dxa')
    _insert_ordered(tblPr, tblW, TBLPR_ORDER)
    if align_center:
        jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); _insert_ordered(tblPr, jc, TBLPR_ORDER)
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side)
        if borders:
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), BORDER)
        else:
            e.set(qn('w:val'), 'nil')
        b.append(e)
    _insert_ordered(tblPr, b, TBLPR_ORDER)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); _insert_ordered(tblPr, lay, TBLPR_ORDER)
    mar = OxmlElement('w:tblCellMar')
    for side, v in zip(('top', 'left', 'bottom', 'right'), cell_mar):
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    _insert_ordered(tblPr, mar, TBLPR_ORDER)
    grid = tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    for w in widths_cm:
        gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(int(w * 567))); grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_cm):
                cell.width = Cm(widths_cm[i])


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:cantSplit'); e.set(qn('w:val'), 'true')
    trPr.insert(0, e)


# ---------------------------------------------------------------- 段落・文字
TAG_RE = re.compile(r'(</?(?:b|r|blue|u|g|w)>)')
HL_RE = re.compile(r'(（要確認）|要確認|○)')


def style_run(run, size, bold=False, color=None, underline=False, font=None):
    run.font.size = Pt(size)
    run.font.bold = bool(bold)
    if underline:
        run.font.underline = True
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    rf = rPr.get_or_add_rFonts()
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), font or FONT)


def add_marked_text(par, text, size=10, bold=False, color=None, font=None):
    """<b><r><blue><u><g><w> の装飾タグを解釈して run を足す"""
    st = {'b': False, 'r': False, 'blue': False, 'u': False, 'g': False, 'w': False}
    for tok in TAG_RE.split(text):
        if not tok:
            continue
        m = re.fullmatch(r'<(/?)(b|r|blue|u|g|w)>', tok)
        if m:
            st[m.group(2)] = (m.group(1) == '')
            continue
        c = color
        b = bold or st['b']
        if st['r']:
            c, b = C_RED, True
        elif st['blue']:
            c, b = C_BLUE, True
        elif st['g']:
            c = C_GRAY
        elif st['w']:
            c, b = 'FFFFFF', True
        # 「要確認」と未設定の目標値「○」は黄色の蛍光ペン（提出前に確認して消すため）
        for part in HL_RE.split(tok):
            if not part:
                continue
            run = par.add_run(part)
            style_run(run, size, b, c, st['u'], font)
            if HL_RE.fullmatch(part):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def fmt_par(par, size=10, exact=True, indent_first=False, hanging=None, align=None, before=0, after=0):
    pf = par.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if exact:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(12 if size >= 10 else (10.5 if size >= 9 else 10))
    else:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if indent_first:
        pf.first_line_indent = Pt(size)
    if hanging is not None:
        pf.left_indent = Pt(hanging); pf.first_line_indent = Pt(-hanging)
    if align == 'center':
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'right':
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == 'left':
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT


def no_grid(p):
    """段落を文書の行グリッド（18pt）に吸着させない。1ptの空段落がグリッド1行分（約20pt）に広がるのを防ぐ"""
    pPr = p._p.get_or_add_pPr()
    for old in pPr.findall(qn('w:snapToGrid')):
        pPr.remove(old)
    e = OxmlElement('w:snapToGrid'); e.set(qn('w:val'), '0')
    # 規格の並び順：snapToGrid は spacing・ind・jc・rPr・sectPr などより前
    for t_ in ('w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
               'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl', 'w:divId',
               'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange'):
        anchor = pPr.find(qn(t_))
        if anchor is not None:
            anchor.addprevious(e)
            return
    pPr.append(e)


def tiny_par(p):
    """空段落を1ptの高さにする（行送り1pt＋段落記号の文字も1pt。LibreOfficeは段落記号の文字の大きさで高さを取るため）"""
    fmt_par(p, 4, exact=True); p.paragraph_format.line_spacing = Pt(1)
    pPr = p._p.get_or_add_pPr()
    for old in pPr.findall(qn('w:rPr')):
        pPr.remove(old)
    rPr = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '2'); rPr.append(sz)
    szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), '2'); rPr.append(szcs)
    anchor = None
    for t_ in ('w:sectPr', 'w:pPrChange'):
        anchor = pPr.find(qn(t_))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(rPr)
    else:
        pPr.append(rPr)
    no_grid(p)


class Container:
    """セル（または本文）に順番にブロックを足す。セル末尾の空段落は再利用する"""

    def __init__(self, cell, width_cm):
        self.cell = cell
        self.width = width_cm
        self._reuse = None   # 直前の add_table が足した空段落

    def para(self):
        if self._reuse is not None:
            p = self._reuse; self._reuse = None
            return p
        cps = self.cell.paragraphs
        if len(self.cell._tc.findall(qn('w:p'))) == 1 and len(self.cell._tc.findall(qn('w:tbl'))) == 0 and not cps[0].text and not cps[0]._p.findall('.//' + qn('w:drawing')):
            if getattr(self, '_first_used', False) is False:
                self._first_used = True
                return cps[0]
        return self.cell.add_paragraph()

    def finish(self):
        """最後に表を足したとき、表の後ろに残る空段落を1ptにする（章末・枠の末尾の余計な空きを防ぐ）"""
        if self._reuse is not None:
            tiny_par(self._reuse)
            self._reuse = None

    def table(self, rows, cols):
        # セルの先頭が表になる場合、セル作成時の空段落が表の上に残って余白になるので取り除く
        tc = self.cell._tc
        lead = None
        if not getattr(self, '_first_used', False):
            ps = tc.findall(qn('w:p'))
            if len(ps) == 1 and not tc.findall(qn('w:tbl')) and not ''.join(ps[0].itertext()).strip() \
                    and not ps[0].findall('.//' + qn('w:drawing')):
                lead = ps[0]
        t = self.cell.add_table(rows=rows, cols=cols)
        if lead is not None:
            # 消すとページ末尾で行を分けられず表示が崩れる（LibreOffice）ため、1ptの段落にして残す
            tiny_par(Paragraph(lead, self.cell))
            self._first_used = True
        # add_table は表の後ろに空段落を足す → 次の段落で再利用
        last = self.cell._tc[-1]
        if last.tag == qn('w:p'):
            self._reuse = Paragraph(last, self.cell)
            fmt_par(self._reuse, size=4, exact=True)
            self._reuse.paragraph_format.line_spacing = Pt(4)
        return t


def keep_group(cont):
    """中身をページの途中で分けない入れ物（枠なし1×1の表・行の分割禁止）。
    見出しや図の題名だけがページ末尾に取り残されるのを防ぐ"""
    t = cont.table(1, 1)
    w = cont.width - 0.15
    set_table_props(t, [w], borders=False, cell_mar=(0, 0, 0, 0))
    cant_split(t.rows[0])
    return Container(t.rows[0].cells[0], w)


def is_small_table(b):
    return b is not None and b['t'] == 'table' and len(b['rows']) <= 5 and not b.get('allow_split')


def render_img(cont, b):
    if b.get('caption'):
        p = cont.para(); fmt_par(p, 9, before=2, align='center'); p.paragraph_format.keep_with_next = True
        add_marked_text(p, b['caption'], 9, True, '1F3864')
    p = cont.para(); fmt_par(p, 10, exact=False, align='center')
    w = min(float(b.get('width', 19.0)), cont.width - 0.2)
    p.add_run().add_picture(os.path.join(HERE, b['src']), width=Cm(w))


def render_blocks(cont, blocks):
    i = 0
    while i < len(blocks):
        b = blocks[i]
        nxt = blocks[i + 1] if i + 1 < len(blocks) else None
        nx2 = blocks[i + 2] if i + 2 < len(blocks) else None
        if b['t'] in ('h2', 'h3') and nxt is not None and nxt['t'] in ('p', 'pn') and nx2 is not None                 and nx2['t'] == 'table' and nx2.get('keep'):
            g = keep_group(cont)          # 小見出し＋説明＋表（"keep"指定）をひとまとめ
            render_one(g, b); render_one(g, nxt); render_one(g, nx2); g.finish()
            i += 3
            continue
        if b['t'] in ('h2', 'h3') and nxt is not None and nxt['t'] == 'img':
            g = keep_group(cont)          # 小見出し＋題名＋図をひとまとめ
            render_one(g, b); render_img(g, nxt); g.finish()
            i += 2
            continue
        if b['t'] in ('h2', 'h3') and nxt is not None and (
                nxt['t'] in ('p', 'pn', 'bullets', 'note') or is_small_table(nxt)):
            g = keep_group(cont)          # 小見出し＋直後の段落（または小さい表）をひとまとめ
            render_one(g, b); render_one(g, nxt); g.finish()
            i += 2
            continue
        if b['t'] in ('p', 'pn') and nxt is not None and is_small_table(nxt) and len(b['text']) <= 70:
            g = keep_group(cont)          # 表の直前の短い前置き（計算式など）＋小さい表をひとまとめ
            render_one(g, b); render_one(g, nxt); g.finish()
            i += 2
            continue
        if is_small_table(b):
            g = keep_group(cont)          # 小さい表はページをまたがせない（結論の行だけが次ページに行かないように）
            render_one(g, b); g.finish()
            i += 1
            continue
        if b['t'] == 'img':
            g = keep_group(cont)          # 題名＋図をひとまとめ
            render_img(g, b); g.finish()
            i += 1
            continue
        render_one(cont, b)
        i += 1
    cont.finish()


def render_one(cont, b):
    if True:
        t = b['t']
        if t == 'h2':
            p = cont.para(); fmt_par(p, 10, before=3); p.paragraph_format.keep_with_next = True
            add_marked_text(p, re.sub(r'^([①-⑳])\s+', r'\1', b['text']), 10, True, C_H2)   # 「① SWOT」→「①SWOT」に統一
        elif t == 'h3':
            p = cont.para(); fmt_par(p, 10, before=1); p.paragraph_format.keep_with_next = True
            add_marked_text(p, b['text'], 10, True, C_H2)
        elif t == 'p':
            p = cont.para(); fmt_par(p, 10, indent_first=True); add_marked_text(p, b['text'], 10)
        elif t == 'pn':
            p = cont.para(); fmt_par(p, 10); add_marked_text(p, re.sub(r'^▶\s+', '▶', b['text']), 10)   # 「▶ 」→「▶」に統一
        elif t == 'note':
            for line in b['text'].split('\n'):
                p = cont.para(); fmt_par(p, 9, hanging=9 if line.startswith('※') else None); add_marked_text(p, line, 9, color=C_NOTE)
        elif t == 'bullets':
            mark = b.get('mark', '・')
            for it in b['items']:
                p = cont.para(); fmt_par(p, 10, hanging=10)
                add_marked_text(p, re.sub(r'^▶\s+', '▶', (mark if not it.startswith(mark) else '') + it), 10)
        elif t == 'space':
            p = cont.para(); fmt_par(p, 4); p.paragraph_format.line_spacing = Pt(4)
        elif t == 'img':
            render_img(cont, b)
        elif t == 'imgrow':
            items = b['items']
            ws = [float(it.get('cell', it['width'] + 0.3)) for it in items]
            tot = sum(ws); maxw = cont.width - 0.15
            if tot > maxw:
                ws = [w * maxw / tot for w in ws]
            t = cont.table(1, len(items))
            set_table_props(t, ws, borders=False, cell_mar=(0, 28, 0, 28))
            for i, it in enumerate(items):
                c = t.rows[0].cells[i]
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
                p = c.paragraphs[0]
                if it.get('caption'):
                    fmt_par(p, 9, align='center'); add_marked_text(p, it['caption'], 9, True, '1F3864')
                    p = c.add_paragraph()
                fmt_par(p, 10, exact=False, align='center')
                p.add_run().add_picture(os.path.join(HERE, it['src']), width=Cm(min(float(it['width']), ws[i] - 0.1)))
        elif t == 'table':
            render_table(cont, b)
        elif t == 'labeled':
            render_labeled(cont, b)
        else:
            raise ValueError('unknown block type: ' + t)


def render_table(cont, b):
    header = b.get('header')
    rows = b['rows']
    ncols = len(b['widths'])
    widths = list(b['widths'])
    total = sum(widths)
    maxw = cont.width - 0.15
    if total > maxw:   # はみ出す場合は比例縮小
        widths = [w * maxw / total for w in widths]
    nrows = len(rows) + (1 if header else 0)
    t = cont.table(nrows, ncols)
    set_table_props(t, widths)
    size = b.get('font', 9)
    align = b.get('align', ['left'] * ncols)
    fills = b.get('fills', {}); row_fills = b.get('row_fills', {}); col_fills = b.get('col_fills', {})
    r0 = 0
    if header:
        hr = t.rows[0]
        set_repeat_header(hr)
        for c, txt in enumerate(header):
            cell = hr.cells[c]
            set_cell_fill(cell, b.get('header_fill', HEADER_FILL))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for k, line in enumerate(str(txt).split('\n')):
                p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                fmt_par(p, size, align='center')
                p.paragraph_format.keep_with_next = True   # 見出し行だけがページ末尾に残らないように
                add_marked_text(p, line, size, True, b.get('header_color', 'FFFFFF'))
        r0 = 1
    for ri, row in enumerate(rows):
        tr = t.rows[ri + r0]
        if not b.get('allow_split'):   # 背の高い行（SWOT等）は分割を許す
            cant_split(tr)
        is_total = b.get('total_row') and ri == len(rows) - 1
        for c in range(ncols):
            txt = row[c] if c < len(row) else ''
            cell = tr.cells[c]
            fill = fills.get(f'{ri},{c}') or row_fills.get(str(ri)) or col_fills.get(str(c)) or ('D9D9D9' if is_total else None)
            if fill:
                set_cell_fill(cell, fill)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP if b.get('valign') == 'top' else WD_CELL_VERTICAL_ALIGNMENT.CENTER
            bold = (b.get('bold_first_col') and c == 0) or is_total
            a_c = align[c] if c < len(align) else 'left'
            nf = b.get('num_font', NUM_FONT) if a_c == 'right' else None   # 右寄せの数字列は等幅書体（ニッポンジーン様と同じ）
            for k, line in enumerate(str(txt).split('\n')):
                p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                fmt_par(p, size, align=a_c)
                add_marked_text(p, line, size, bold, font=nf)
    for m in b.get('merge', []):
        r1, c1, r2, c2 = m
        a = t.cell(r1 + r0, c1); z = t.cell(r2 + r0, c2)
        keep = a.text
        merged = a.merge(z)
        # 結合で増えた空段落を整理
        ps = merged.paragraphs
        if any(p.text.strip() for p in ps):
            for p in ps:
                if not p.text.strip() and len(merged.paragraphs) > 1:
                    p._p.getparent().remove(p._p)
    return t


def render_labeled(cont, b):
    lw = float(b.get('label_width', 2.4))
    items = b['items']
    t = cont.table(len(items), 2)
    rw = cont.width - 0.15 - lw
    set_table_props(t, [lw, rw])
    for i, it in enumerate(items):
        lc, rc = t.rows[i].cells
        set_cell_fill(lc, 'DBE5F1')
        lc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for k, line in enumerate(it['label'].split('\n')):
            p = lc.paragraphs[0] if k == 0 else lc.add_paragraph()
            fmt_par(p, 10, align='center'); add_marked_text(p, line, 10, True, C_H2)
        sub = Container(rc, rw - 0.2)
        render_blocks(sub, it['blocks'])


def clear_cell(cell):
    tc = cell._tc
    for ch in list(tc):
        if ch.tag != qn('w:tcPr'):
            tc.remove(ch)
    p = OxmlElement('w:p'); tc.append(p)


def drop_row_height(table, also_cantsplit=True):
    for tr in table._tbl.findall(qn('w:tr')):
        trPr = tr.find(qn('w:trPr'))
        if trPr is None:
            continue
        for tag in ('w:trHeight',) + (('w:cantSplit',) if also_cantsplit else ()):
            for e in trPr.findall(qn(tag)):
                trPr.remove(e)


def remove_par(p_el):
    p_el.getparent().remove(p_el)


# ---------------------------------------------------------------- 前段（（1）〜（4））
FRONT = json.load(open(os.path.join(HERE, 'front.json'), encoding='utf-8'))


def build(out, template):
    doc = Document(template)
    sec = doc.sections[0]
    sec.top_margin = Emu(421 * 635); sec.right_margin = Emu(397 * 635); sec.bottom_margin = Emu(426 * 635); sec.left_margin = Emu(284 * 635)
    body = doc.element.body
    # 様式由来の浮動図形（「参考様式」テキストボックス、説明文を囲む赤枠など）を除去
    MC = '{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent'
    for el in list(body.iterchildren()):
        if el.tag == qn('w:p'):
            for mc in list(el.iter(MC)):
                mc.getparent().remove(mc)
            for pict in list(el.iter(qn('w:pict'))):
                pict.getparent().remove(pict)
    children = list(body.iterchildren())
    tables = [c for c in children if c.tag == qn('w:tbl')]
    TW = 19.6
    # タイトル
    p0 = Paragraph(children[0], doc)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    fmt_par(p0, 11, align='center'); add_marked_text(p0, '事業計画書', 11, True)
    # 見出し段落（（１）〜（５））のフォント統一。字下げは（１）の設定（表の枠とそろう）に全見出しを合わせる
    ind_ref = None
    for el in children:
        if el.tag == qn('w:p') and re.match(r'\s*（１）', ''.join(x.text or '' for x in el.iter(qn('w:t')))):
            pPr = el.find(qn('w:pPr'))
            ind_ref = pPr.find(qn('w:ind')) if pPr is not None else None
            break
    for el in children:
        if el.tag == qn('w:p'):
            par = Paragraph(el, doc)
            txt = par.text
            if re.match(r'\s*（[１２３４５]）', txt) or txt.startswith('その1'):
                for r in list(par.runs):
                    r._r.getparent().remove(r._r)
                fmt_par(par, 10, before=2); add_marked_text(par, txt.strip(), 10, True)
                # 様式の見出しごとに字下げが違い左端がずれるので、（１）と同じ字下げにそろえる
                pPr = par._p.get_or_add_pPr()
                for old in pPr.findall(qn('w:ind')):
                    pPr.remove(old)
                if ind_ref is not None:
                    new_ind = copy.deepcopy(ind_ref)
                    # w:ind は pPr の中で spacing の後・jc の前に置く（規格の並び順）
                    nxt_tags =('w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')
                    anchor = None
                    for t_ in nxt_tags:
                        anchor = pPr.find(qn(t_))
                        if anchor is not None:
                            break
                    if anchor is not None:
                        anchor.addprevious(new_ind)
                    else:
                        pPr.append(new_ind)
            elif txt.startswith('※補助金交付候補者') or txt.startswith('※事業計画に沿って') or txt.startswith('ただし、補助事業として'):
                for r in list(par.runs):
                    r._r.getparent().remove(r._r)
                fmt_par(par, 8); add_marked_text(par, txt.strip(), 8, color='595959')
    # 様式の記入説明（削除）
    for el in children:
        if el.tag == qn('w:p'):
            txt = ''.join(x.text or '' for x in el.iter(qn('w:t')))
            if txt.startswith(('※以下のフォーマット等は', '　※本事業計画書（その１・その２）', '外部支援者や認定支援機関等', '（システムに入力いただく', '支援者は、真に')):
                remove_par(el)
    # （4）の表と（5）見出しの間の空段落を1つに減らし、（5）は新しいページから始める
    #   （本文の表の冒頭だけが1ページ目の末尾に残るのを防ぐ）
    blanks = []
    el = tables[3].getnext()
    while el is not None and not (el.tag == qn('w:p') and ''.join(x.text or '' for x in el.iter(qn('w:t'))).strip()):
        if el.tag == qn('w:p') and not ''.join(x.text or '' for x in el.iter(qn('w:t'))).strip():
            blanks.append(el)
        el = el.getnext()
    for b_el in blanks[1:]:
        remove_par(b_el)
    if el is not None and ''.join(x.text or '' for x in el.iter(qn('w:t'))).strip().startswith('（５）'):
        Paragraph(el, doc).paragraph_format.page_break_before = True
    # （1）事業者情報
    t1 = Table(tables[0], doc)
    set_table_props(t1, [4.0, TW - 4.0])
    for r, (label, val) in enumerate(FRONT['jigyosha']):
        lc, vc = t1.rows[r].cells[0], t1.rows[r].cells[1]
        if label is not None:
            clear_cell(lc); p = lc.paragraphs[0]; fmt_par(p, 9); add_marked_text(p, label, 9)
        if val is not None:
            clear_cell(vc); p = vc.paragraphs[0]; fmt_par(p, 9); add_marked_text(p, val, 9)
    # （2）（3）
    for ti, key in ((1, 'keikakumei'), (2, 'gaiyo')):
        t = Table(tables[ti], doc); set_table_props(t, [TW]); drop_row_height(t)
        c = t.rows[0].cells[0]; clear_cell(c); p = c.paragraphs[0]; fmt_par(p, 10); add_marked_text(p, FRONT[key], 10, key == 'keikakumei')
    # （4）補助経費の一覧＋図面
    t4 = Table(tables[3], doc); set_table_props(t4, [TW]); drop_row_height(t4)
    c4 = t4.rows[0].cells[0]; clear_cell(c4)
    cont4 = Container(c4, TW - 0.2)
    render_blocks(cont4, FRONT['keihi_blocks'])
    # （5）本文：12行×1列
    tm = Table(tables[4], doc); set_table_props(tm, [TW], cell_mar=(28, 85, 28, 85)); drop_row_height(tm)
    files = {int(re.search(r'row(\d+)', f).group(1)): f for f in glob.glob(os.path.join(HERE, 'chapters', 'row*.json'))}
    for r in range(12):
        cell = tm.rows[r].cells[0]
        if r not in files:
            continue
        ch = json.load(open(files[r], encoding='utf-8'))
        clear_cell(cell)
        cont = Container(cell, TW - 0.3)
        blocks = ch['blocks']
        # 章見出し＋最初の段落をひとまとめ（見出しだけがページ末尾に残らないように）
        g = keep_group(cont)
        p = g.para(); fmt_par(p, 10.5, before=2); p.paragraph_format.line_spacing = Pt(14); p.paragraph_format.keep_with_next = True
        add_marked_text(p, ch['title'], 10.5, True, C_CH)
        k = 1 if blocks and blocks[0]['t'] in ('p', 'pn') else 0
        render_blocks(g, blocks[:k])
        render_blocks(cont, blocks[k:])
    # 本文の表とスケジュールのページの間の段落（横向きへの切替え＝セクション区切りを持つ）は、
    # 様式では高さ約1cm。本文の最終ページからこれだけがはみ出して白紙のページができるので、1ptにする
    sep = tables[4].getnext()
    if sep is not None and sep.tag == qn('w:p') and not ''.join(x.text or '' for x in sep.iter(qn('w:t'))).strip():
        sp = Paragraph(sep, doc)
        sp.paragraph_format.space_before = Pt(0); sp.paragraph_format.space_after = Pt(0)
        sp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; sp.paragraph_format.line_spacing = Pt(1)
        no_grid(sp)
        spc = sep.find(qn('w:pPr')).find(qn('w:spacing'))
        for a in ('w:afterLines', 'w:beforeLines', 'w:afterAutospacing', 'w:beforeAutospacing'):
            if spc is not None and spc.get(qn(a)) is not None:
                del spc.attrib[qn(a)]
    # スケジュールのページ（空白）
    ts = Table(tables[5], doc)
    cs = ts.rows[0].cells[0]; clear_cell(cs)
    p = cs.paragraphs[0]; fmt_par(p, 10.5); add_marked_text(p, '補助事業のスケジュール', 10.5, True, C_CH)
    for el in list(body.iterchildren()):
        if el.tag == qn('w:p'):
            txt = ''.join(x.text or '' for x in el.iter(qn('w:t')))
            if txt.startswith('事業計画は事業者ごとの決算期'):
                for tnode in el.iter(qn('w:t')):
                    tnode.text = ''
    doc.save(out)
    print('saved:', out)


if __name__ == '__main__':
    out = sys.argv[1]
    tpl = sys.argv[sys.argv.index('--template') + 1] if '--template' in sys.argv else os.path.join(HERE, 'template.docx')
    build(out, tpl)
