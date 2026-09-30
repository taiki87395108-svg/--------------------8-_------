# -*- coding: utf-8 -*-
"""v2 → v3 生成スクリプト（python-pptx）。v2 は触らず、v3 を別名で保存する。"""
import sys, json, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from lxml import etree

BASE = r'C:\Users\taiki\Dropbox\中小企業診断士\★省力化投資補助金\★【クロードコード用】省力化投資補助金第8回_トンボ飲料様\作成中事業計画書'
SRC = os.path.join(BASE, '★0930Gライン_レイアウト工程ビフォーアフター_トンボ飲料様_v2.pptx')
DST = os.path.join(BASE, '★0930Gライン_レイアウト工程ビフォーアフター_トンボ飲料様_v3.pptx')
IMG_JSON = sys.argv[1] if len(sys.argv) > 1 else None   # 画像リスト（json: [{path, caption, credit}]）

FONT = 'BIZ UDPGothic'
TEAL = '0F4C5C'; ORANGE = 'F26B1D'; DKOR = 'B84A0A'; GRAY = '5B6770'; FOOT = '8A9BA8'; INK = '1F2A30'
PANEL = 'F7F9FA'; ROW = 'EEF4F5'; ROWLN = 'C9D3D8'; HI = 'FDE8DC'; BARLN = 'D5DEE2'; WHITE = 'FFFFFF'
POUCH = 'D9DEE3'; POUCHLN = '7F8C96'; LIGHT = 'E6EEF1'


def rgb(h):
    return RGBColor.from_string(h)


def set_run_font(r, size, bold=False, color=INK):
    r.font.size = Pt(size)
    r.font.bold = bool(bold)
    r.font.color.rgb = rgb(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set('typeface', FONT)


def fill_tf(tf, paras, size=9, bold=False, color=INK, align=None, anchor='ctr', spacing=None):
    """paras: str | list[ str | list[(text, size, bold, color)] ]
    既存図形では先頭段落の段落書式（箇条書き等）を2段落目以降にも引き継ぐ"""
    import copy
    pPr0 = tf.paragraphs[0]._p.find(qn('a:pPr'))
    pPr_copy = copy.deepcopy(pPr0) if pPr0 is not None else None
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = {'ctr': MSO_ANCHOR.MIDDLE, 't': MSO_ANCHOR.TOP, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i > 0 and pPr_copy is not None and p._p.find(qn('a:pPr')) is None:
            p._p.insert(0, copy.deepcopy(pPr_copy))
        if align is not None:
            p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if spacing:
            p.line_spacing = spacing
        segs = [(para, size, bold, color)] if isinstance(para, str) else para
        for seg in segs:
            t = seg[0]
            sz = seg[1] if len(seg) > 1 and seg[1] is not None else size
            b = seg[2] if len(seg) > 2 and seg[2] is not None else bold
            c = seg[3] if len(seg) > 3 and seg[3] is not None else color
            r = p.add_run()
            r.text = t
            set_run_font(r, sz, b, c)


def add_rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh


def add_text(slide, x, y, w, h, paras, size=9, bold=False, color=INK, align='l', anchor='ctr', spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.auto_size = MSO_AUTO_SIZE.NONE
    fill_tf(tf, paras, size, bold, color, align, anchor, spacing)
    return tb


def add_picture_fit(slide, path, x, y, w, h):
    from PIL import Image
    im = Image.open(path)
    iw, ih = im.size
    r = min(w / iw, h / ih)
    pw, ph = iw * r, ih * r
    return slide.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def find_shape(slide, sid=None, contains=None):
    for sh in slide.shapes:
        if sid is not None and sh.shape_id == sid:
            return sh
        if contains and sh.has_text_frame and contains in sh.text_frame.text:
            return sh
    raise KeyError((sid, contains))


def set_shape_text(slide, sid, paras, **kw):
    sh = find_shape(slide, sid)
    fill_tf(sh.text_frame, paras, **kw)
    return sh


def set_cell_text(cell, paras):
    """表セルの先頭runの書式（サイズ・太字・色）と配置を引き継いで文字を差し替える"""
    tf = cell.text_frame
    p0 = tf.paragraphs[0]
    align = p0.alignment
    r0 = p0.runs[0] if p0.runs else None
    size = r0.font.size.pt if r0 is not None and r0.font.size else 9
    bold = bool(r0.font.bold) if r0 is not None else False
    color = INK
    if r0 is not None:
        c = r0._r.find(qn('a:rPr'))
        c = c.find(qn('a:solidFill')) if c is not None else None
        c = c.find(qn('a:srgbClr')) if c is not None else None
        if c is not None:
            color = c.get('val')
    tf.clear()
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        segs = [(para, size, bold, color)] if isinstance(para, str) else para
        for seg in segs:
            r = p.add_run(); r.text = seg[0]
            set_run_font(r, seg[1] if len(seg) > 1 and seg[1] else size,
                         seg[2] if len(seg) > 2 and seg[2] is not None else bold,
                         seg[3] if len(seg) > 3 and seg[3] else color)


def header(slide, title, subtitle, pageno):
    add_rect(slide, 0, 0, 10, 0.75, fill=TEAL, line=TEAL, lw=1)
    add_rect(slide, 0.35, 0.20, 0.12, 0.35, fill=ORANGE, line=ORANGE, lw=1)
    add_text(slide, 0.58, 0.08, 9.10, 0.60, title, size=22, bold=True, color=WHITE)
    add_text(slide, 0.35, 0.80, 9.30, 0.34, subtitle, size=11, color=GRAY)
    add_text(slide, 0.35, 5.28, 7.0, 0.25, '株式会社トンボ飲料様　Gライン パウチ自動供給装置の導入', size=9, color=FOOT)
    add_text(slide, 9.10, 5.28, 0.55, 0.25, str(pageno), size=9, color=FOOT, align='r')


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    bg = s.background
    bg.fill.solid(); bg.fill.fore_color.rgb = rgb(WHITE)
    return s


def move_slide(prs, slide, new_index):
    lst = prs.slides._sldIdLst
    items = list(lst)
    target = None
    for it in items:
        if prs.slides.part.related_part(it.rId) is slide.part:  # noqa
            target = it
            break
    if target is None:
        for it in items:
            if prs.part.related_part(it.rId) is slide.part:
                target = it
                break
    lst.remove(target)
    lst.insert(new_index, target)


def kpi_card(slide, x, y, big, sub):
    add_rect(slide, x, y, 2.98, 0.74, fill=ROW, line=None)
    add_text(slide, x + 0.15, y + 0.02, 2.70, 0.38, big, size=17, bold=True, color=DKOR)
    add_text(slide, x + 0.15, y + 0.38, 2.70, 0.34, sub, size=8, color=GRAY)


def draw_pouch(slide, x, y, w=0.46, h=0.62, spout=False, cap=False, weld=False, filled=False):
    top = 0.16 if spout else 0.0
    add_rect(slide, x, y + top, w, h - top, fill=('F6E3C8' if filled else POUCH), line=POUCHLN, lw=0.75,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    if spout:
        sw = 0.13
        add_rect(slide, x + w / 2 - sw / 2, y + 0.04, sw, 0.16, fill=WHITE, line=TEAL, lw=1)
        if weld:
            add_rect(slide, x + w / 2 - 0.16, y + top - 0.02, 0.32, 0.05, fill=ORANGE, line=ORANGE, lw=0.25)
        if cap:
            add_rect(slide, x + w / 2 - 0.11, y - 0.02, 0.22, 0.08, fill=ORANGE, line=ORANGE, lw=0.25)


# ---------------------------------------------------------------- main
prs = Presentation(SRC)
S = list(prs.slides)   # 0:title 1:全体 2:レイアウト 3:役割 4:工程BA 5:シフト 6:設備

# ===== slide 2（全体の流れ）=====
s = S[1]
sh = set_shape_text(s, 5, '製造ラインは A・D（ドリンク類：セレブレ・シャンメリー等）／F（スティックゼリー）／E・G（パウチゼリー）の5本。'
                          '本事業はGライン全体が対象で、装置は製袋室に導入し、投入の自動化で安定稼働と人員の再配置につなげます。',
                    size=9.5, color=GRAY)
sh.top = Inches(0.78); sh.height = Inches(0.38)
sh = set_shape_text(s, 23, [
    [('人手 投入担当（②〜⑦）', 8, True, TEAL)],
    [('開梱・ほぐし→踏み台からマガジンへ手投入（3分詰めて5分で空、の繰り返し）', 8, False, INK)],
    [('人手 オペレーター（⑧）', 8, True, TEAL)],
    [('製袋機の運転・監視', 8, False, INK)],
    [('自動：スパウト（ストロー）圧着→検査', 8, False, GRAY)]], anchor='t')
sh.top = Inches(2.17); sh.height = Inches(1.05)
set_shape_text(s, 31, [[('充填室・殺菌室（1F）', 10, True, WHITE)]])
sh = set_shape_text(s, 32, [
    [('人手 充填担当1名（⑨）', 8, True, TEAL)],
    [('G仕込室で調合・殺菌した液を充填→キャップ→殺菌室（リフトスチーマー）で容器ごと殺菌', 8, False, INK)],
    [('自動：キャップ供給（ホッパローダ）', 8, False, GRAY)]], anchor='t')
sh.top = Inches(2.17); sh.height = Inches(1.05)
add_rect(s, 4.60, 1.47, 2.40, 0.25, fill=ROW, line=ROWLN, lw=0.75)
add_text(s, 4.66, 1.47, 2.30, 0.25, '※ G仕込室（1F）から調合液を送る', size=7.5, color=GRAY)
set_shape_text(s, 15, [[('資材保管室（1F）', 9, True, WHITE)]])
set_shape_text(s, 39, [
    [('人手 包装担当2〜3名（⑩）', 9, True, TEAL)],
    [('包装・箱詰め→1階へ下ろして製品倉庫・出荷', 9, False, INK)]], anchor='t')
set_shape_text(s, 12, '効果：投入担当の確保に左右されず17.5時間稼働を維持。創出した時間（約13.3時間／日）は多能工化・品質管理へ再配分',
               size=9, bold=True, color=TEAL)
set_shape_text(s, 53, ['稼働 8:30〜26:00（17.5時間）', '早番 6:30〜11:00・中番 11:00〜19:00', '遅番 19:00〜26:00（夜勤）',
                       '遅番の投入は派遣で賄えず社員がカバー出勤'], size=8.5)
# 表紙の出所
set_shape_text(S[0], 7, '出所：フジシール提案資料（2025/12/22・2026/3/25）、御見積書（2026/9/24）、ヒアリングリスト（v5）、配置図面（4号棟1F・2F）、'
                        '社内打合せ（2026/9/30）、事業計画書（その3）別紙1', size=9, color=FOOT)
set_shape_text(s, 57, ['1品種の日 15万袋／日（月11日）', '2品種の日 10万袋／日（月11日）', '製袋機の能力 150袋／分',
                       'Gライン売上 約14億円（全社売上比 約13%）'], size=8.5)

# ===== slide 3（レイアウト）=====
s = S[2]
set_shape_text(s, 29, ['投入担当（パート）1名 ②〜⑦', '手でほぐしてマガジンへ投入', '3分詰めて5分で空、の繰り返し',
                       '1日約131回、離れられない'], size=8)
set_shape_text(s, 37, 'オペレーター1名 ⑧', size=8.5, bold=True)
set_shape_text(s, 86, [[('② パレット詰め', 7.5, True, DKOR)], [('約2分／枚', 7.5, False, INK)], [('早番・中番', 7.5, False, INK)]])
sh = set_shape_text(s, 91, [[('オペレーター1名 ⑤', 8.5, True, INK)], [('（遅番はパレット入替③を兼務）', 8.5, True, INK)]])
sh.width = Inches(2.0)
set_shape_text(s, 80, ['ロボットアームが', 'マガジンへ供給'], size=7.5)
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == '装填':
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                r.font.size = Pt(7.5)
set_shape_text(s, 104, [[('変わる点：', 9.5, True, DKOR),
                         ('踏み台での手投入（3分詰めて5分で空、の繰り返し）がなくなり、マガジン脇に自動供給装置を設置。'
                          '人は専用パレットへ詰めるだけになり、投入担当を常時置く必要がなくなる。製袋機より後ろの工程は変更なし。', 9.5, False, INK)]])

# ===== slide 4（各機械の役割）=====
s = S[3]
tbl = find_shape(s, 8).table
set_cell_text(tbl.cell(1, 4), ['自動供給に合わせて', '部品改造・設計変更'])
set_cell_text(tbl.cell(8, 3), ['1枚1,200袋。型替えで2種のパウチに兼用（約1〜1.5分／枚）', '装置内25枚（5段×4列＋予備1列）＋予備65枚'])
set_cell_text(tbl.cell(8, 4), ['詰め作業に使用', '（約2分／枚）'])
set_cell_text(tbl.cell(10, 1), '充填・殺菌設備')
set_cell_text(tbl.cell(10, 2), '既設（1F）')
set_cell_text(tbl.cell(10, 3), ['G仕込室で調合液を仕込み・殺菌 → 充填室で充填・キャップ', '→ 殺菌室で容器ごと殺菌（リフトスチーマー）'])
set_shape_text(s, 5, '記号a〜kは前ページのレイアウト図と対応。j・kは製袋室の外（図にはありません）。', size=11, color=GRAY)

# ===== slide 5（工程ビフォーアフター）: 行を作り直す =====
s = S[4]
for sh in list(s.shapes):
    if sh.shape_id >= 8:
        sh._element.getparent().remove(sh._element)
set_shape_text(s, 5, '①〜の番号と数値は別紙1と同じ。オレンジ＝左は無くなる作業、右は新たな作業。下3段は前後で変わらない作業。',
               size=11, color=GRAY)


def row(slide, x, y, label, value, note, hi=False):
    add_rect(slide, x, y, 4.35, 0.28, fill=HI if hi else ROW, line=ORANGE if hi else ROWLN, lw=1.5 if hi else 0.75)
    add_text(slide, x + 0.08, y, 2.15, 0.28, label, size=9, bold=True, color=INK)
    add_text(slide, x + 2.20, y, 0.78, 0.28, value, size=10.5, bold=True, color=DKOR if hi else TEAL)
    add_text(slide, x + 3.05, y, 1.27, 0.28, note, size=7.5, color=GRAY)


# 左：現状
add_rect(s, 0.35, 1.20, 4.55, 3.62, fill=PANEL, line=TEAL, lw=1.25)
add_rect(s, 0.35, 1.20, 0.90, 0.30, fill=GRAY, line=GRAY, lw=1)
add_text(s, 0.35, 1.20, 0.90, 0.30, '現状', size=12, bold=True, color=WHITE, align='c')
add_text(s, 1.35, 1.20, 3.50, 0.30, '投入担当の作業 1,050分／日（①〜⑦の合計）', size=9.5, bold=True, color=TEAL)
L = [('① 空パウチの搬入', '20分', ['数箱ずつ5〜6回', '（変わらない）'], False),
     ('② 段ボールの開梱・束出し', '43分', ['43箱×1分', '（詰め作業に含める）'], True),
     ('③ 空段ボールの片付け', '22分', ['43箱×0.5分', '（詰め作業に含める）'], True),
     ('④ パウチのほぐし', '563分', ['無くなる', '（投入の合間に継続）'], True),
     ('⑤ 踏み台の昇降・残量確認', '79分', ['無くなる', '（0.6分×131回）'], True),
     ('⑥ マガジンへの手投入', '315分', ['無くなる', '（2.4分×131回）'], True),
     ('⑦ 品種切替時のパウチ切替', '8分', ['無くなる（月平均）', '2品種の日に15分'], True),
     ('⑧ 製袋機の運転・監視', '1,050分', 'オペレーター1名', False),
     ('⑨ 充填（充填室）', '1,050分', '1名', False),
     ('⑩ 包装・箱詰め', '2,625分', '2〜3名（平均2.5名）', False)]
for i, (a, b, c, hi) in enumerate(L):
    row(s, 0.45, 1.56 + 0.32 * i, a, b, c, hi)
# 右：導入後
add_rect(s, 5.10, 1.20, 4.55, 3.62, fill=PANEL, line=TEAL, lw=1.25)
add_rect(s, 5.10, 1.20, 0.90, 0.30, fill=ORANGE, line=ORANGE, lw=1)
add_text(s, 5.10, 1.20, 0.90, 0.30, '導入後', size=12, bold=True, color=WHITE, align='c')
add_text(s, 6.10, 1.20, 3.50, 0.30, '投入に代わる作業 251分／日（①〜④。うち新規231分）', size=9.5, bold=True, color=TEAL)
R = [('① 空パウチの搬入', '20分', '変化なし', False),
     ('② 専用パレットへの詰め', '209分', ['新規（約2分／枚）', '250／168分の平均'], True),
     ('③ パレットの装填・入替', '15分', ['新規（月平均）', '約3分×約2時間に1回'], True),
     ('④ 専用パレットの型替え', '7分', ['新規（月平均7分）', '2品種の日の半数で28分'], True)]
for i, (a, b, c, hi) in enumerate(R):
    row(s, 5.20, 1.56 + 0.32 * i, a, b, c, hi)
add_rect(s, 5.20, 2.84, 4.35, 0.92, fill=WHITE, line=ROWLN, lw=0.75)
add_text(s, 5.30, 2.84, 4.15, 0.92, [[('（該当する作業なし）', 8, False, FOOT)],
                                     [('②③の開梱・片付けは詰め作業②に含めて計上', 8, False, GRAY)],
                                     [('④〜⑦はロボットが供給するため無くなる', 8, False, GRAY)]], align='c')
R2 = [('⑤ 製袋機の運転・監視', '1,050分', ['供給はロボット', '（人手0分）']),
      ('⑥ 充填（充填室）', '1,050分', '変化なし'),
      ('⑦ 包装・箱詰め', '2,625分', '変化なし')]
for i, (a, b, c) in enumerate(R2):
    row(s, 5.20, 3.80 + 0.32 * i, a, b, c, False)
add_rect(s, 0.35, 4.90, 9.30, 0.34, fill=ROW, line=BARLN, lw=1)
add_text(s, 0.45, 4.90, 9.10, 0.34, [[('別紙1の計算：', 10, True, TEAL),
                                     ('削減 1,030分／日（17.2時間） − 発生 231分／日（3.9時間） ＝ 799分／日（13.3時間）　', 10, False, INK),
                                     ('省力化指数 0.78', 10, True, DKOR), ('（＝799÷1,030）', 10, False, INK)]])

# ===== slide 6（シフト別）=====
s = S[5]
set_shape_text(s, 9, '1,050分 → 251分', size=17, bold=True, color=DKOR)
set_shape_text(s, 10, '投入担当の作業→導入後の関連作業（月平均。搬入20分は変わらず）', size=8, color=GRAY)
set_shape_text(s, 12, '約799分／日 創出', size=17, bold=True, color=DKOR)
set_shape_text(s, 13, '≒13.3時間／日。主に中番と遅番で創出', size=8, color=GRAY)
set_shape_text(s, 15, '省力化指数 0.78', size=17, bold=True, color=DKOR)
set_shape_text(s, 16, '（1,030−231）÷1,030。別紙1の自動計算と同じ', size=8, color=GRAY)
tbl = find_shape(s, 17).table
rows = [
    ['早番', ['6:30〜11:00', '（投入は8:30から）'], '投入パート1名', '同じ方が詰め作業へ転換', '150', '150／150', '0／0（業務転換）'],
    ['中番', '11:00〜19:00', '投入パート1名', '早番の残りを詰め、以降は多能工化へ', '480', '100／18', '380／462'],
    ['遅番（夜勤）', '19:00〜26:00', ['派遣1名', '（社員がカバーも）'], '投入専任なし。入替はオペレーター兼務', '420', '0／0', '420／420'],
    ['搬入（変わらない）', '全シフト', '上3行の内数（20分）', '詰め担当が同じ回数を運ぶ（20分）', '（内数）', '20／20', '−20／−20'],
    ['装填・入替', '約2時間に1回×約3分', 'なし', ['新規。日中は詰め担当、', '遅番はオペレーターが兼務'], '0', '18／12', '−18／−12'],
    ['型替え', 'パウチ形状が変わる日', 'なし', ['新規。機内25枚×約1.1分≒28分', '（2品種の日の半数で発生→月平均7分）'], '0', '0／14', '0／−14'],
    ['合計', '8:30〜26:00（17.5時間）', '延べ3名／日', '詰めは早番・中番のみ', '1,050', ['288／214', '（平均251）'], ['762／836', '（平均799）']],
]
for ri, vals in enumerate(rows, start=1):
    for ci, v in enumerate(vals):
        set_cell_text(tbl.cell(ri, ci), v)

# ===== slide 7（設備・見積）=====
s = S[6]
tbl = find_shape(s, 14).table
set_cell_text(tbl.cell(5, 1), ['5段×4列＝20パレット＝24,000袋', '（約2時間分：フジシール資料）'])
set_cell_text(tbl.cell(6, 1), ['1枚1,200袋、詰め 約2分／枚', '型替え（2種兼用）約1〜1.5分／枚'])
set_cell_text(tbl.cell(9, 1), '発注から7〜9か月（メーカーの受注状況次第）')
tbl2 = find_shape(s, 15).table
set_cell_text(tbl2.cell(1, 0), 'パウチ自動供給装置（ロボットアーム式）')
set_cell_text(tbl2.cell(6, 0), '現地改造工事費（約11日の工事・調整）')
for sid in (8, 9, 11, 12):          # 左右パネルと縦帯を少し下へ延ばす
    find_shape(s, sid).height = Inches(3.92)
sh = set_shape_text(s, 16, [
    [('見積条件：', 7.5, True, GRAY), ('消費税別・概算金額。ユーティリティ工事は別途。車上渡し条件。', 7.5, False, GRAY)],
    [('トンボ飲料様の手配範囲：', 7.5, True, GRAY), ('荷降ろし・設置（アンカー固定まで）、搬入経路上の機器の一時解体、試運転用の製品・資材。', 7.5, False, GRAY)],
    [('補助申請：', 7.5, True, GRAY), ('補助対象経費92,000,000円×補助率1/2＝補助金額46,000,000円。見積は値引き行の削除と名称の修正を依頼中（総額は不変）。', 7.5, False, GRAY)]],
    anchor='t')
sh.top = Inches(4.40); sh.height = Inches(0.70)

# ===== 新スライドA：パウチとスパウト =====
imgs = json.load(open(IMG_JSON, encoding='utf-8')) if IMG_JSON and os.path.exists(IMG_JSON) else []
sa = new_slide(prs)
header(sa, '2. パウチとスパウト（ストロー）とは ― 製袋機で行う工程',
       'Gラインで作る「パウチゼリー」の容器の説明。本事業で自動化するのは、この容器の元になる空パウチを製袋機へ投入する作業です。', 3)
# 左：写真
add_rect(sa, 0.35, 1.20, 4.55, 3.62, fill=PANEL, line=TEAL, lw=1.25)
add_rect(sa, 0.35, 1.20, 1.60, 0.30, fill=GRAY, line=GRAY, lw=1)
add_text(sa, 0.35, 1.20, 1.60, 0.30, '製品と容器の写真', size=11, bold=True, color=WHITE, align='c')
if imgs:
    # 2列グリッド（最大4枚）。各セル：写真（白地）＋説明文
    gap = 0.12
    cols = 2
    cw = (4.35 - gap * (cols - 1)) / cols
    rows_n = (len(imgs) + cols - 1) // cols
    ch = (2.96 - gap * (rows_n - 1)) / rows_n          # 1.60〜4.56 を使う
    img_h = ch - 0.36
    for i, im in enumerate(imgs):
        r, c = divmod(i, cols)
        x = 0.45 + c * (cw + gap)
        y = 1.60 + r * (ch + gap)
        add_rect(sa, x, y, cw, img_h, fill=WHITE, line=ROWLN, lw=0.75)
        add_picture_fit(sa, im['path'], x + 0.04, y + 0.04, cw - 0.08, img_h - 0.08)
        add_text(sa, x, y + img_h + 0.02, cw, 0.34, im['caption'], size=8, color=INK, anchor='t')
    seen = []
    for im in imgs:
        if im.get('credit') and im['credit'] not in seen:
            seen.append(im['credit'])
    add_text(sa, 0.45, 4.60, 4.35, 0.18, '出所：' + '、'.join(seen) + ' の各Webサイト（製品・容器の紹介ページ）', size=7, color=FOOT)
else:
    add_text(sa, 0.45, 1.60, 4.35, 3.10, '（写真を挿入）', size=10, color=FOOT, align='c')
# 右：製袋機の工程
add_rect(sa, 5.10, 1.20, 4.55, 3.62, fill=PANEL, line=TEAL, lw=1.25)
add_rect(sa, 5.10, 1.20, 2.30, 0.30, fill=ORANGE, line=ORANGE, lw=1)
add_text(sa, 5.10, 1.20, 2.30, 0.30, '製袋機の工程（スパウトの圧着）', size=11, bold=True, color=WHITE, align='c')
steps = [
    (dict(spout=False), [('① 空パウチ（平らな袋）', 9, True, TEAL)],
     '1束50袋・段ボールで納品。投入担当が手でマガジン（製袋機の投入口）へ入れる ← 本事業で自動化する作業'),
    (dict(spout=True, weld=True), [('② 製袋機（既設SPM機）でスパウトを熱圧着', 9, True, TEAL)],
     '袋の口にストロー状の注ぎ口「スパウト」を熱で溶着する（自動、150袋／分）'),
    (dict(spout=True, weld=True), [('③ リークチェッカで検査', 9, True, TEAL)],
     '圧着部からの漏れがないか検査（自動）→ コンベアで充填室へ'),
    (dict(spout=True, weld=True, cap=True, filled=True), [('④ 充填室で充填・キャップ', 9, True, TEAL)],
     'ゼリーを充填してキャップを付け、殺菌室で容器ごと殺菌 → 2階で包装・箱詰め'),
]
y0 = 1.60
for i, (opt, ttl, body) in enumerate(steps):
    y = y0 + i * 0.78
    add_rect(sa, 5.22, y, 4.32, 0.70, fill=WHITE, line=ROWLN, lw=0.75)
    draw_pouch(sa, 5.40, y + 0.05, 0.46, 0.60, **opt)
    add_text(sa, 6.05, y + 0.03, 3.42, 0.24, [ttl], size=9, bold=True, color=TEAL)
    add_text(sa, 6.05, y + 0.27, 3.42, 0.42, body, size=7.5, color=INK, anchor='t')
    if i < len(steps) - 1:
        add_rect(sa, 5.55, y + 0.70, 0.16, 0.08, fill=TEAL, line=TEAL, lw=0.25, shape=MSO_SHAPE.DOWN_ARROW)
add_text(sa, 5.22, 4.66, 4.32, 0.14, 'オレンジ＝圧着部・キャップ。図は概念図（寸法は実物と異なります）', size=7, color=FOOT)
add_rect(sa, 0.35, 4.90, 9.30, 0.34, fill=ROW, line=BARLN, lw=1)
add_text(sa, 0.45, 4.90, 9.10, 0.34, [[('用語：', 8.5, True, TEAL),
                                      ('パウチ＝アルミ等を貼り合わせた柔らかい袋。スパウト＝飲み口になるストロー状の注ぎ口（キャップ付き）。'
                                       '製袋機＝空パウチにスパウトを圧着する機械（既設SPM機）。マガジン＝製袋機の投入口', 8.5, False, INK)]])

# ===== 新スライドB：投入作業のサイクル =====
sb = new_slide(prs)
header(sb, '5. 投入作業のサイクル（現状 vs 導入後）',
       '現状は「3分詰めて5分で空」の繰り返しで持ち場を離れられない。導入後は約2時間に1回の入替（約3分）だけになる（同じ4時間分で比較）', 6)
find_shape(sb, contains='同じ4時間分で比較').text_frame.paragraphs[0].runs[0].font.size = Pt(10)
X0, W = 0.35, 9.30
scale = W / 240.0
# 目盛り
for t in range(0, 241, 60):
    lab = '開始' if t == 0 else f'{t // 60}時間後'
    add_text(sb, X0 + t * scale - (0.5 if t == 240 else 0.0), 1.18, 0.5, 0.16, lab, size=7.5, color=FOOT,
             align='r' if t == 240 else 'l')
    add_rect(sb, X0 + t * scale - 0.005, 1.34, 0.01, 0.10, fill=FOOT, line=FOOT, lw=0.25)
# 現状の帯
add_text(sb, X0, 1.44, 3.0, 0.22, [[('現状：手投入', 9.5, True, GRAY)]])
add_text(sb, X0 + 3.0, 1.44, 6.3, 0.22, '約8分サイクル（投入3分＋空になるまで5分）が4時間で30回、1日では約131回', size=8, color=GRAY, align='r')
add_rect(sb, X0, 1.68, W, 0.42, fill=LIGHT, line=ROWLN, lw=0.5)
for k in range(30):
    add_rect(sb, X0 + k * 8 * scale, 1.68, 3 * scale, 0.42, fill=ORANGE, line=None)
add_text(sb, X0 + 0.02, 2.12, 6.0, 0.16, '■ 投入（約3分）　□ 空になるまで（約5分：次の束のほぐし・開梱に追われる）', size=8, color=GRAY)
# 拡大図（1サイクル）
add_text(sb, X0, 2.36, 2.0, 0.20, '1サイクルの拡大（約8分）', size=8.5, bold=True, color=TEAL)
cx, cy, cw, ch = X0, 2.58, 4.40, 0.46
add_rect(sb, cx, cy, cw * 3 / 8, ch, fill=ORANGE, line=None)
add_rect(sb, cx + cw * 3 / 8, cy, cw * 5 / 8, ch, fill=LIGHT, line=ROWLN, lw=0.5)
add_text(sb, cx + 0.05, cy, cw * 3 / 8 - 0.1, ch, [[('投入 約3分', 8.5, True, WHITE)], [('踏み台に上がり束を詰める', 8, False, WHITE)]], align='c')
add_text(sb, cx + cw * 3 / 8 + 0.05, cy, cw * 5 / 8 - 0.1, ch,
         [[('約5分で空になる', 8.5, True, INK)], [('合間も次の束のほぐし・開梱で手が空かない', 8, False, INK)]], align='c')
add_text(sb, 4.95, 2.36, 4.70, 0.70,
         [[('1日の稼働1,050分 ÷ 8分 ≒ 約131回の投入。', 8.5, True, INK)],
          [('投入担当（パート1名）はシフト中ずっと持ち場に拘束され、早番150分・中番480分・遅番420分＝1,050分／日。'
            '遅番は派遣が確保できず、日勤社員のカバー出勤が常態化', 8.5, False, INK)]], anchor='t')
# 導入後の帯
add_text(sb, X0, 3.18, 3.0, 0.22, [[('導入後：自動供給', 9.5, True, DKOR)]])
add_text(sb, X0 + 3.0, 3.18, 6.3, 0.22, '装置に約2時間分（パレット20枚＝24,000袋、フジシール資料）を装填。入替は約3分×4時間で2回', size=8, color=GRAY, align='r')
add_rect(sb, X0, 3.42, W, 0.42, fill=LIGHT, line=ROWLN, lw=0.5)
for t0 in (117, 237):
    add_rect(sb, X0 + t0 * scale, 3.42, 3 * scale, 0.42, fill=ORANGE, line=None)
for t0 in (0, 120):
    add_text(sb, X0 + t0 * scale + 0.1, 3.42, 117 * scale - 0.2, 0.42,
             [[('ロボットが自動供給 約2時間', 9, True, TEAL)], [('パウチ詰めは早番・中番が自分のペースでまとめて実施', 8, False, GRAY)]], align='c')
add_text(sb, X0 + 0.02, 3.86, 6.0, 0.16, '■ パレットの入替（約3分。遅番は製袋室オペレーターが兼務）　□ 人手不要', size=8, color=GRAY)
add_text(sb, X0, 4.06, 4.55, 0.40,
         [[('パレット詰め（約2分／枚）', 8.5, True, TEAL), ('は早番・中番がまとめて行う → 1品種の日250分、2品種の日168分（月平均209分）', 8.5, False, INK)]], anchor='t')
add_text(sb, 5.10, 4.06, 4.55, 0.40,
         [[('入替（約3分）', 8.5, True, TEAL), ('は生産量÷24,000袋で1品種の日6回＝18分、2品種の日4回＝12分。遅番はオペレーターが兼務し、投入担当が不要に', 8.5, False, INK)]], anchor='t')
kpi_card(sb, 0.35, 4.50, '1,050分 → 251分／日', '投入の拘束時間 → 導入後の関連作業（搬入20＋詰め209＋入替15＋型替え7、月平均）')
kpi_card(sb, 3.51, 4.50, '約799分／日 創出', '≒13.3時間。主に中番・遅番で創出')
kpi_card(sb, 6.67, 4.50, '省力化指数 0.78', '（1,030−231）÷1,030。別紙1の自動計算と同じ')

# ===== 並び替えと番号 =====
move_slide(prs, sa, 2)   # 全体の次
move_slide(prs, sb, 5)   # 各機械の役割の次
titles = ['1. Gライン全体の流れと人の配置（現状）', None, '3. 製袋室のレイアウト ビフォー・アフター（デフォルメ図）', '4. 各機械の役割', None,
          '6. パウチ供給工程のビフォー・アフター（1日の人手作業時間）', '7. 人の配置・担当人数・工数の変化（シフト別）', '8. 導入する設備の概要と見積の構成']
for i, sl in enumerate(list(prs.slides)[1:], start=2):
    t = titles[i - 2]
    if t is None:
        continue
    set_shape_text(sl, 4, t, size=22, bold=True, color=WHITE)
    set_shape_text(sl, 7, str(i), size=9, color=FOOT, align='r')

prs.save(DST)
print('saved', DST, 'slides', len(prs.slides))
